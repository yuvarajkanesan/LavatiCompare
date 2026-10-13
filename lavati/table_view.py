"""CSV / Excel / HTML table comparison by key column."""
import os

from PySide6.QtCore import QRectF, Qt, Signal
from PySide6.QtGui import QColor, QFontMetrics, QPainter
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog, QFrame,
                               QHBoxLayout, QHeaderView, QLabel, QMessageBox, QPushButton, QScrollArea,
                               QStyledItemDelegate, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget)

from . import engine, theme
from .widgets import sty, toast, Segmented, card

MAX_ROWS = 1500


def kind_bg(k):
    return QColor({"left": theme.ROW_LEFT, "right": theme.ROW_RIGHT}[k])


def kind_fg(k):
    return QColor({"left": theme.BLUE, "right": theme.PURPLE}[k])


def _is_num(s):
    try:
        float(s.replace(",", ""))
        return True
    except ValueError:
        return False


class CellDelegate(QStyledItemDelegate):
    def paint(self, p, opt, idx):
        data = idx.data(Qt.UserRole)
        if data is None:
            return super().paint(p, opt, idx)
        left, right, changed, kind, numeric, key = data
        r = opt.rect
        p.save()
        p.setRenderHint(QPainter.Antialiasing)
        if kind in ("left", "right"):
            p.fillRect(r, kind_bg(kind))
        elif changed:
            p.fillRect(r, QColor(theme.ROW_CHANGED))
        else:
            p.fillRect(r, QColor(theme.PANEL))
        p.setPen(QColor(theme.LINE_SOFT))
        p.drawLine(r.bottomLeft(), r.bottomRight())
        font = theme.mono(10) if (numeric or key) else theme.ui(10)
        p.setFont(font)
        fm = QFontMetrics(font)
        align = Qt.AlignRight if numeric else Qt.AlignLeft
        area = r.adjusted(10, 0, -10, 0)
        if changed:
            new_font = theme.mono(10, True) if numeric else theme.ui(10, True)
            nfm = QFontMetrics(new_font)
            nw = nfm.horizontalAdvance(right) + 16
            ow = fm.horizontalAdvance(left)
            total = ow + 8 + nw
            x = area.right() - total if numeric else area.left()
            f = p.font()
            f.setStrikeOut(True)
            p.setFont(f)
            p.setPen(QColor(theme.GRAY))
            p.drawText(QRectF(x, r.top(), ow, r.height()), Qt.AlignVCenter | Qt.AlignLeft, left)
            chip = QRectF(x + ow + 8, r.top() + 6, nw, r.height() - 12)
            p.setPen(Qt.NoPen)
            p.setBrush(QColor(theme.CELL_NEW))
            p.drawRoundedRect(chip, 5, 5)
            p.setFont(new_font)
            p.setPen(QColor(theme.CELL_NEW_TEXT))
            p.drawText(chip, Qt.AlignCenter, right)
        else:
            text = left if kind != "right" else right
            p.setPen((kind_fg(kind) if kind in ("left", "right") else QColor(theme.TEXT)))
            p.drawText(area, Qt.AlignVCenter | align, text)
        p.restore()


class SheetSummary(QDialog):
    """Compares every sheet pair that has the same name (Excel) and lists the differences per sheet."""

    def __init__(self, left, right, key_choice, opts, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Sheet summary")
        self.resize(720, 420)
        lay = QVBoxLayout(self)
        t = QTableWidget(0, 6)
        t.setHorizontalHeaderLabels(["Sheet", "Rows left", "Rows right", "Changed", "Only left", "Only right"])
        t.verticalHeader().hide()
        t.setEditTriggers(QTableWidget.NoEditTriggers)
        t.horizontalHeader().setStretchLastSection(True)
        t.setColumnWidth(0, 200)
        try:
            ls, rs = engine.list_table_sources(left), engine.list_table_sources(right)
        except Exception as e:
            ls = rs = []
            QMessageBox.warning(self, "Sheets", str(e))
        names = [n for n in ls if n in rs] + [n for n in ls if n not in rs] + [n for n in rs if n not in ls]
        t.setRowCount(len(names))
        for i, n in enumerate(names):
            vals = [n]
            if n in ls and n in rs:
                try:
                    L, R = engine.load_table(left, n), engine.load_table(right, n)
                    common = [h for h in L[0] if h in R[0]]
                    key = key_choice if key_choice in common else (common[0] if common else None)
                    if key is None:
                        vals += [str(len(L[1])), str(len(R[1])), "no common column", "", ""]
                    else:
                        res = engine.compare_tables(L, R, key, set(common) - {key}, **opts)
                        c = res.counts
                        vals += [str(len(L[1])), str(len(R[1])), str(c["changed"]), str(c["left"]), str(c["right"])]
                except Exception as e:
                    vals += ["", "", f"error: {e}", "", ""]
            else:
                vals += ["", "", "sheet exists on one side only", "", ""]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                if c >= 3 and v.isdigit() and int(v) > 0:
                    it.setForeground(QColor(theme.RED))
                t.setItem(i, c, it)
        lay.addWidget(t)
        bb = QDialogButtonBox(QDialogButtonBox.Close)
        bb.rejected.connect(self.reject)
        bb.button(QDialogButtonBox.Close).clicked.connect(self.accept)
        lay.addWidget(bb)


class TableCompareView(QWidget):
    resultChanged = Signal(str)

    def __init__(self, cfg, left, right, parent=None):
        super().__init__(parent)
        self.cfg = cfg
        self.left, self.right = left, right
        self.mode = 1
        self.result = None
        self.sources = {"L": [], "R": []}
        self.sheet = {"L": None, "R": None}
        self.L = self.R = ([], [])
        self.key_name = ""
        self._loading = False
        self._ui()
        self.reload()

    # ------------------------------------------------------------- loading
    def _read_sources(self):
        for side, path in (("L", self.left), ("R", self.right)):
            try:
                self.sources[side] = engine.list_table_sources(path)
            except Exception:
                self.sources[side] = []
            if self.sheet[side] not in self.sources[side]:
                self.sheet[side] = self.sources[side][0] if self.sources[side] else None
        if self.same_sheet.isChecked() and self.sheet["L"] in self.sources["R"]:
            self.sheet["R"] = self.sheet["L"]

    def reload(self):
        self._read_sources()
        self._load_tables()

    def _load_tables(self):
        errors = []
        for side, path in (("L", self.left), ("R", self.right)):
            try:
                t = engine.load_table(path, self.sheet[side])
            except Exception as e:
                t = ([], [])
                errors.append(f"{os.path.basename(path)}: {e}")
            if side == "L":
                self.L = t
            else:
                self.R = t
        if errors:
            QMessageBox.warning(self, "Load error", "\n".join(errors))
        self.headers = list(self.L[0]) + [h for h in self.R[0] if h not in self.L[0]]
        self._build_side()
        self.compare()

    # ------------------------------------------------------------------ UI
    def _ui(self):
        root = QHBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        self.side = QFrame()
        self.side.setFixedWidth(360)
        sty(self.side, f"background:{theme.SIDEBAR}; border-right:1px solid {theme.BORDER};")
        self.side_lay = QVBoxLayout(self.side)
        self.side_lay.setContentsMargins(22, 22, 22, 22)
        self.side_lay.setSpacing(8)
        self.same_sheet = QCheckBox("Pick the same sheet / table on both sides")
        self.same_sheet.setChecked(True)
        self.same_sheet.setStyleSheet("border:none;")
        sc = QScrollArea()
        sc.setWidget(self.side)
        sc.setWidgetResizable(True)
        sc.setFixedWidth(361)
        sc.setFrameShape(QFrame.NoFrame)
        root.addWidget(sc)

        main = QWidget()
        ml = QVBoxLayout(main)
        ml.setContentsMargins(28, 24, 28, 18)
        ml.setSpacing(14)
        cards = QHBoxLayout()
        cards.setSpacing(16)
        self.stat = {}
        for key, title, col in (("changed", "Changed rows", theme.AMBER), ("right", "Only in right (new)", theme.PURPLE),
                                ("left", "Only in left (removed)", theme.BLUE), ("same", "Identical rows", theme.TEXT)):
            c = card()
            cl = QVBoxLayout(c)
            cl.setContentsMargins(20, 14, 20, 14)
            cl.setSpacing(2)
            t = QLabel(title)
            t.setStyleSheet(f"color:{theme.MUTED}; border:none;")
            v = QLabel("0")
            v.setStyleSheet(f"color:{col}; font-size:22pt; font-weight:700; border:none;")
            cl.addWidget(t)
            cl.addWidget(v)
            self.stat[key] = v
            cards.addWidget(c, 1)
        ml.addLayout(cards)
        r2 = QHBoxLayout()
        self.seg = Segmented(["All", "Differences"], current=1)
        self.seg.changed.connect(lambda i: (setattr(self, "mode", i), self.render()))
        r2.addWidget(self.seg)
        r2.addStretch(1)
        self.b_sheets = QPushButton("  All sheets summary…")
        self.b_sheets.setIcon(theme.icon("table", theme.TEXT, 18))
        self.b_sheets.setProperty("plain", True)
        self.b_sheets.clicked.connect(self.sheet_summary)
        ex = QPushButton("  Export diff (.xlsx)")
        ex.setIcon(theme.icon("download", theme.TEXT, 18))
        ex.setProperty("plain", True)
        ex.clicked.connect(self.export)
        r2.addWidget(self.b_sheets)
        r2.addWidget(ex)
        ml.addLayout(r2)
        self.table = QTableWidget()
        self.table.setItemDelegate(CellDelegate(self.table))
        self.table.verticalHeader().hide()
        self.table.setShowGrid(False)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setSelectionMode(QTableWidget.NoSelection)
        self.table.setStyleSheet(f"QTableWidget {{ border:1px solid {theme.BORDER}; border-radius:10px; }}")
        ml.addWidget(self.table, 1)
        self.foot = QLabel()
        self.foot.setStyleSheet(f"color:{theme.MUTED};")
        ml.addWidget(self.foot)
        root.addWidget(main, 1)

    def _clear_side(self):
        def clear(layout):
            while layout.count():
                it = layout.takeAt(0)
                w = it.widget()
                if w is not None:
                    w.setParent(None)
                    if w is not self.same_sheet:
                        w.deleteLater()
                elif it.layout() is not None:
                    clear(it.layout())
        clear(self.side_lay)

    def _build_side(self):
        self._loading = True
        self._clear_side()
        sl = self.side_lay

        def head(t):
            l = QLabel(t)
            l.setStyleSheet(f"color:{theme.MUTED}; font-weight:700; font-size:9pt; letter-spacing:1px; border:none;")
            return l

        def sub(t):
            l = QLabel(t)
            l.setStyleSheet("font-weight:700; font-size:11pt; border:none; margin-top:10px;")
            return l
        sl.addWidget(head("FILES"))
        self.sheet_combos = {}
        for tag, path, tbl in (("L", self.left, self.L), ("R", self.right, self.R)):
            l = QLabel(f"{tag} · {os.path.basename(path)} ({len(tbl[1]):,} rows)")
            l.setFont(theme.mono(9))
            l.setWordWrap(True)
            l.setStyleSheet(f"background:{theme.PANEL}; border:1px solid {theme.BORDER}; border-radius:6px; padding:8px 10px;")
            l.setToolTip(path)
            sl.addWidget(l)
            if len(self.sources[tag]) > 1 or (self.sources[tag] and path.lower().endswith((".html", ".htm"))):
                cb = QComboBox()
                cb.addItems(self.sources[tag])
                cb.setCurrentText(self.sheet[tag] or "")
                cb.currentTextChanged.connect(lambda t_, tag=tag: self._sheet_changed(tag, t_))
                self.sheet_combos[tag] = cb
                sl.addWidget(cb)
        if any(len(v) > 1 for v in self.sources.values()):
            sl.addWidget(self.same_sheet)
        sl.addWidget(sub("Key column"))
        self.key = QComboBox()
        common = [h for h in self.L[0] if h in self.R[0]] or self.headers
        self.key.addItems(common)
        if self.key_name in common:
            self.key.setCurrentText(self.key_name)
        sl.addWidget(self.key)
        sl.addWidget(sub("Columns to compare"))
        self.col_checks = {}
        for h in self.headers:
            cb = QCheckBox(h)
            cb.setStyleSheet("border:none;")
            cb.setChecked(True)
            cb.toggled.connect(lambda on, c=cb, name=h: c.setText(name if on else f"{name}  (ignored)"))
            self.col_checks[h] = cb
            sl.addWidget(cb)
        sl.addWidget(sub("Options"))
        self.o_case = QCheckBox("Ignore case")
        self.o_trim = QCheckBox("Trim whitespace")
        self.o_text = QCheckBox("Numbers as text")
        self.o_rel = QCheckBox("Tolerance is a percentage")
        for c in (self.o_case, self.o_trim):
            c.setChecked(True)
        for c in (self.o_case, self.o_trim, self.o_text, self.o_rel):
            c.setStyleSheet("border:none;")
            sl.addWidget(c)
        row = QHBoxLayout()
        lab = QLabel("Numeric tolerance ±")
        lab.setStyleSheet("border:none;")
        self.tol = QDoubleSpinBox()
        self.tol.setDecimals(4)
        self.tol.setButtonSymbols(QDoubleSpinBox.NoButtons)
        self.tol.setRange(0, 1e9)
        self.tol.setValue(0.01)
        row.addWidget(lab)
        row.addWidget(self.tol, 1)
        sl.addLayout(row)
        sl.addSpacing(12)
        btn = QPushButton("Re-compare")
        btn.setProperty("primary", True)
        btn.setMinimumHeight(42)
        btn.clicked.connect(self.compare)
        sl.addWidget(btn)
        sl.addStretch(1)
        self.b_sheets.setVisible(any(len(v) > 1 for v in self.sources.values()))
        self._loading = False

    def _sheet_changed(self, side, name):
        if self._loading:
            return
        self.sheet[side] = name
        if self.same_sheet.isChecked():
            other = "R" if side == "L" else "L"
            if name in self.sources[other]:
                self.sheet[other] = name
        self._load_tables()

    # -------------------------------------------------------------
    def _opts(self):
        return dict(ignore_case=self.o_case.isChecked(), trim=self.o_trim.isChecked(),
                    num_as_text=self.o_text.isChecked(), tol=self.tol.value(), relative=self.o_rel.isChecked())

    def compare(self):
        if not self.headers:
            self.result = engine.TableResult([], [], {"changed": 0, "left": 0, "right": 0, "same": 0})
            self.key_name = ""
            for v in self.stat.values():
                v.setText("0")
            self.table.setRowCount(0)
            return
        key = self.key.currentText()
        cols = {h for h, cb in self.col_checks.items() if cb.isChecked() and h != key}
        self.result = engine.compare_tables(self.L, self.R, key, cols, **self._opts())
        self.key_name = key
        c = self.result.counts
        for k, v in self.stat.items():
            v.setText(f"{c[k]:,}")
        self.resultChanged.emit(f"{c['changed'] + c['left'] + c['right']} row diffs")
        self.render()

    def sheet_summary(self):
        SheetSummary(self.left, self.right, self.key.currentText(), self._opts(), self).exec()

    def render(self):
        res = self.result
        if res is None or not res.headers:
            return
        heads = res.headers
        rows = [r for r in res.rows if self.mode == 0 or r["kind"] != "same"]
        total = len(rows)
        rows = rows[:MAX_ROWS]
        numeric = []
        for ci, h in enumerate(heads):
            vals = [c[ci][0] or c[ci][1] for c in (r["cells"] for r in res.rows[:200])]
            vals = [v for v in vals if v.strip()]
            numeric.append(bool(vals) and all(_is_num(v) for v in vals))
        self.table.clear()
        self.table.setColumnCount(len(heads) + 1)
        self.table.setRowCount(len(rows))
        labels = [""] + [h + ("   KEY" if h == self.key_name else "") for h in heads]
        self.table.setHorizontalHeaderLabels(labels)
        for c, h in enumerate(heads):
            self.table.horizontalHeaderItem(c + 1).setTextAlignment(
                (Qt.AlignRight if numeric[c] else Qt.AlignLeft) | Qt.AlignVCenter)
        hh = self.table.horizontalHeader()
        hh.setHighlightSections(False)
        hh.setSectionResizeMode(0, QHeaderView.Fixed)
        self.table.setColumnWidth(0, 44)
        for c in range(1, len(heads) + 1):
            hh.setSectionResizeMode(c, QHeaderView.Stretch)
        glyph = {"changed": ("", theme.AMBER), "left": ("-", theme.BLUE), "right": ("+", theme.PURPLE), "same": ("", theme.GRAY)}
        for ri, r in enumerate(rows):
            self.table.setRowHeight(ri, 40)
            g, gc = glyph[r["kind"]]
            gi = QTableWidgetItem(g)
            gi.setForeground(QColor(gc))
            gi.setTextAlignment(Qt.AlignCenter)
            f = gi.font()
            f.setBold(True)
            gi.setFont(f)
            if r["kind"] in ("left", "right"):
                gi.setBackground(kind_bg(r["kind"]))
            if r["kind"] == "changed":
                gi.setIcon(theme.status_icon("changed"))
            self.table.setItem(ri, 0, gi)
            for ci, (a, b, ch) in enumerate(r["cells"]):
                it = QTableWidgetItem()
                it.setData(Qt.UserRole, (a, b, ch, r["kind"], numeric[ci], heads[ci] == self.key_name))
                self.table.setItem(ri, ci + 1, it)
        diffs = res.counts["changed"] + res.counts["left"] + res.counts["right"]
        shown_diffs = sum(1 for r in rows if r["kind"] != "same")
        where = ""
        if self.sheet["L"] or self.sheet["R"]:
            where = f" · {self.sheet['L'] or ''}" + (f" vs {self.sheet['R']}" if self.sheet["R"] != self.sheet["L"] else "")
        self.foot.setText((f"Showing {shown_diffs:,} of {diffs:,} differences" if self.mode == 1 else f"Showing {len(rows):,} of {total:,} rows")
                          + f" · rows matched on {self.key_name}{where}" + (f" · first {MAX_ROWS:,} rows only" if total > MAX_ROWS else ""))

    def export(self):
        path, _ = QFileDialog.getSaveFileName(self, "Export diff", "diff.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Export", "openpyxl is not installed.")
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Diff"
        ws.append(["Status"] + self.result.headers)
        fills = {"changed": "FFF3D6", "left": "E3ECFB", "right": "EDE7FB"}
        for r in self.result.rows:
            if r["kind"] == "same":
                continue
            vals = []
            for a, b, ch in r["cells"]:
                vals.append(f"{a} → {b}" if ch else (a or b))
            ws.append([{"changed": "Changed", "left": "Removed", "right": "New"}[r["kind"]]] + vals)
            row = ws.max_row
            for ci, (a, b, ch) in enumerate(r["cells"]):
                cell = ws.cell(row=row, column=ci + 2)
                if ch:
                    cell.fill = PatternFill("solid", fgColor="FBE6BD")
                    cell.font = Font(bold=True)
                elif r["kind"] != "changed":
                    cell.fill = PatternFill("solid", fgColor=fills[r["kind"]])
        for c in ws.columns:
            ws.column_dimensions[c[0].column_letter].width = min(40, max(len(str(x.value or "")) for x in c) + 3)
        wb.save(path)
        toast(f"Exported {os.path.basename(path)}")

    def has_unsaved(self):
        return False

    def swap_sides(self):
        self.left, self.right = self.right, self.left
        self.sheet["L"], self.sheet["R"] = self.sheet["R"], self.sheet["L"]
        self.reload()
