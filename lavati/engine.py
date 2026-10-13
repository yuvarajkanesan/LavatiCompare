"""Comparison engines. No Qt imports here."""
import csv
import difflib
import fnmatch
import os
import re
from dataclasses import dataclass, field

from . import vfs

TEXT_EXTS = {
    ".txt", ".md", ".py", ".kt", ".java", ".xml", ".gradle", ".json", ".yml", ".yaml", ".html",
    ".htm", ".css", ".js", ".ts", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".sh", ".bat",
    ".ps1", ".ini", ".cfg", ".toml", ".properties", ".pro", ".sql", ".log", ".kts", ".swift",
    ".rb", ".php", ".csv", ".tsv", ".gitignore", ".mk", ".cmake", ".rc", ".vue", ".jsx", ".tsx",
    ".reg", ".inf", ".conf", ".svg", ".tex", ".lua", ".pl", ".r", ".scss", ".less", ".gradle",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".ico", ".tif", ".tiff"}
TABLE_EXTS = {".csv", ".tsv", ".xlsx", ".xlsm"}
HTML_EXTS = {".html", ".htm"}
MEDIA_EXTS = {".mp3", ".flac", ".m4a", ".mp4", ".aac", ".m4b", ".m4v"}
DOC_EXTS = {".docx", ".pdf", ".odt", ".rtf", ".pptx", ".odp", ".docm"}
PE_EXTS = {".exe", ".dll", ".sys", ".ocx", ".scr", ".cpl", ".mui", ".efi"}

STATUSES = ("same", "changed", "newer", "older", "left", "right")
DIFF_STATUSES = ("changed", "newer", "older", "left", "right")
BOTH_SIDE_DIFF = ("changed", "newer", "older")      # exists on both sides but differs


class Cancelled(Exception):
    pass


class Ctx:
    """Shared state so a worker thread can report progress and be cancelled."""

    def __init__(self):
        self.count = 0
        self.cancel = False


# =============================================================== text rules
WS_MODES = ("none", "edges", "changes", "all")
UNIMPORTANT_PRESETS = {
    "C-style line comments (// ...)": r"//.*$",
    "Script comments (# ...)": r"#.*$",
    "SQL / Lua comments (-- ...)": r"--.*$",
    "Block comments on one line (/* ... */)": r"/\*.*?\*/",
    "XML / HTML comments (<!-- ... -->)": r"<!--.*?-->",
    "ISO timestamps (2024-01-31 12:00:00)": r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}([.,]\d+)?)?(Z|[+-]\d{2}:?\d{2})?",
    "Times (12:34:56)": r"\b\d{1,2}:\d{2}:\d{2}\b",
    "Version-control keywords ($Id: ... $)": r"\$(Id|Date|Revision|Author|Header)[^$]*\$",
}


@dataclass
class TextRules:
    ws: str = "none"                       # none | edges | changes | all
    ignore_case: bool = False
    ignore_eol: bool = True
    unimportant: list = field(default_factory=list)      # regex strings: matches are removed from the key
    replacements: list = field(default_factory=list)     # [(regex, replacement)] applied first
    _comp: tuple = field(default=None, repr=False, compare=False)

    @classmethod
    def from_config(cls, cfg):
        ws = "changes" if cfg["ignore_whitespace"] else "none"
        return cls(ws=ws, ignore_case=bool(cfg["ignore_case_text"]), ignore_eol=bool(cfg["ignore_eol"]),
                   unimportant=list(cfg["unimportant"]), replacements=[tuple(x) for x in cfg["replacements"]])

    def copy(self, **kw):
        d = dict(ws=self.ws, ignore_case=self.ignore_case, ignore_eol=self.ignore_eol,
                 unimportant=list(self.unimportant), replacements=list(self.replacements))
        d.update(kw)
        return TextRules(**d)

    def is_plain(self):
        return (self.ws == "none" and not self.ignore_case and not self.unimportant and not self.replacements)

    def _compile(self):
        if self._comp is None:
            reps, unimp = [], []
            for pat, rep in self.replacements:
                try:
                    reps.append((re.compile(pat), rep))
                except re.error:
                    pass
            for pat in self.unimportant:
                try:
                    unimp.append(re.compile(pat))
                except re.error:
                    pass
            self._comp = (reps, unimp)
        return self._comp

    def invalid_patterns(self):
        bad = []
        for pat in list(self.unimportant) + [p for p, _ in self.replacements]:
            try:
                re.compile(pat)
            except re.error as e:
                bad.append(f"{pat}: {e}")
        return bad

    def key(self, s):
        if self.is_plain():
            return s
        reps, unimp = self._compile()
        for rx, rep in reps:
            s = rx.sub(rep, s)
        for rx in unimp:
            s = rx.sub("", s)
        if self.ws == "edges":
            s = s.strip()
        elif self.ws == "changes":
            s = " ".join(s.split())
        elif self.ws == "all":
            s = "".join(s.split())
        if self.ignore_case:
            s = s.lower()
        return s

    def keys(self, lines, eol_tag=""):
        if self.is_plain() and not eol_tag:
            return lines
        if eol_tag:
            return [self.key(x) + eol_tag for x in lines]
        return [self.key(x) for x in lines]


# =============================================================== folders
@dataclass
class Node:
    name: str
    rel: str
    is_dir: bool
    left: dict = None     # {"size","mtime","path"} or None
    right: dict = None
    status: str = "same"  # same | changed | newer | older | left | right
    children: list = field(default_factory=list)
    depth: int = 0
    rname: str = ""       # right-hand name when paired by an alignment override
    archive: bool = False
    mismatch: bool = False
    approx: bool = False  # decided from metadata only (snapshot / no checksum available)
    lfs: object = None
    rfs: object = None
    renamed_from: str = ""

    def right_name(self):
        return self.rname or self.name


@dataclass
class CompareOptions:
    method: str = "size_time"      # size_time | crc | binary | rules
    time_tol: float = 2.0
    ignore_case_names: bool = False
    patterns: list = field(default_factory=list)      # exclude (ignore) patterns
    include: list = field(default_factory=list)       # file include patterns; empty = everything
    archives: bool = False
    aligns: list = field(default_factory=list)        # [[left_rel, right_rel], ...]
    rules: TextRules = None                           # for method "rules"

    @classmethod
    def from_config(cls, cfg, **kw):
        o = cls(method=cfg["compare_method"], time_tol=float(cfg["time_tolerance"]),
                ignore_case_names=bool(cfg["ignore_case_names"]), patterns=list(cfg["ignore_patterns"]),
                archives=bool(cfg["archives"]), rules=TextRules.from_config(cfg))
        for k, v in kw.items():
            setattr(o, k, v)
        return o

    def to_dict(self):
        r = self.rules
        return dict(method=self.method, time_tol=self.time_tol, ignore_case_names=self.ignore_case_names,
                    patterns=list(self.patterns), include=list(self.include), archives=self.archives,
                    aligns=[list(a) for a in self.aligns],
                    rules=dict(ws=r.ws, ignore_case=r.ignore_case, ignore_eol=r.ignore_eol,
                               unimportant=list(r.unimportant), replacements=[list(x) for x in r.replacements]) if r else None)

    def apply_dict(self, d):
        for k in ("method", "time_tol", "ignore_case_names", "patterns", "include", "archives", "aligns"):
            if k in d:
                setattr(self, k, d[k])
        if d.get("rules"):
            rd = d["rules"]
            self.rules = TextRules(ws=rd.get("ws", "none"), ignore_case=rd.get("ignore_case", False),
                                   ignore_eol=rd.get("ignore_eol", True), unimportant=list(rd.get("unimportant", [])),
                                   replacements=[tuple(x) for x in rd.get("replacements", [])])
        return self


def is_ignored(name, is_dir, patterns):
    if name == ".lavati-backup":
        return True
    for p in patterns:
        p = p.strip()
        if not p:
            continue
        dir_only = p.endswith("/") or p.endswith("\\")
        p = p.rstrip("/\\")
        if dir_only and not is_dir:
            continue
        if fnmatch.fnmatch(name.lower(), p.lower()):
            return True
    return False


def matches_include(name, patterns):
    pats = [p.strip() for p in patterns if p.strip()]
    if not pats:
        return True
    n = name.lower()
    return any(fnmatch.fnmatch(n, p.lower()) for p in pats)


def _info(fs, path, ent, is_dir):
    return {"size": 0 if is_dir else ent.size, "mtime": ent.mtime, "path": path}


def _list(fs, path, opts):
    """name-key -> (Entry, full path). Applies ignore + include filters."""
    out = {}
    if fs is None or path is None:
        return out
    for e in fs.scandir(path):
        if is_ignored(e.name, e.is_dir, opts.patterns):
            continue
        if not e.is_dir and not matches_include(e.name, opts.include):
            continue
        key = e.name.lower() if opts.ignore_case_names else e.name
        while key in out:           # same name differing only in case on one side
            key += "\0"
        out[key] = (e, fs.join(path, e.name))
    return out


def looks_binary(chunk):
    if chunk.startswith((b"\xff\xfe", b"\xfe\xff", b"\xef\xbb\xbf")):
        return False
    return b"\x00" in chunk


def _binary_equal(lfs, lp, rfs, rp, ctx):
    with lfs.open(lp) as a, rfs.open(rp) as b:
        while True:
            x, y = a.read(vfs.CHUNK), b.read(vfs.CHUNK)
            if x != y:
                return False
            if not x:
                return True
            if ctx is not None and ctx.cancel:
                raise Cancelled()


def _read_all(fs, path, limit=64 * 1024 * 1024):
    with fs.open(path) as f:
        data = f.read(limit + 1)
    return data if len(data) <= limit else None


def _rules_equal(lfs, lp, rfs, rp, opts, ctx):
    a, b = _read_all(lfs, lp), _read_all(rfs, rp)
    if a is None or b is None or looks_binary(a[:8192]) or looks_binary(b[:8192]):
        return _binary_equal(lfs, lp, rfs, rp, ctx)
    ta, ea, _ = decode_bytes(a)
    tb, eb, _ = decode_bytes(b)
    rules = opts.rules or TextRules()
    la, lb = split_lines(ta), split_lines(tb)
    tag_a = tag_b = ""
    if not rules.ignore_eol:
        tag_a, tag_b = "\0" + detect_eol(ta)[0], "\0" + detect_eol(tb)[0]
    return rules.keys(la, tag_a) == rules.keys(lb, tag_b)


def files_differ(lfs, lp, le, rfs, rp, re_, opts, ctx=None):
    """Return (differs, approx). `approx` means only metadata could be used."""
    method = opts.method
    tol = opts.time_tol
    if (lfs.cheap_algo == "git" or rfs.cheap_algo == "git") and method != "rules":
        a, b = lfs.checksum(lp, "git", ctx), rfs.checksum(rp, "git", ctx)
        if a is not None and b is not None:
            return a != b, False
    if method == "size_time" or not (lfs.has_content and rfs.has_content):
        if method != "size_time":
            ca, cb = lfs.checksum(lp, "crc32", ctx), rfs.checksum(rp, "crc32", ctx)
            if ca is not None and cb is not None:
                return ca != cb, False
        same = le.size == re_.size and abs(le.mtime - re_.mtime) <= tol
        return (not same), (method != "size_time")
    if method != "rules" and le.size != re_.size:
        return True, False
    if method == "crc":
        algo = lfs.cheap_algo or rfs.cheap_algo      # e.g. S3 supplies MD5 ETags without a download
        a = lfs.checksum(lp, algo, ctx) if algo else None
        b = rfs.checksum(rp, algo, ctx) if algo else None
        if a is None or b is None:
            a, b = lfs.checksum(lp, "crc32", ctx), rfs.checksum(rp, "crc32", ctx)
        return a != b, False
    if method == "binary":
        if le.size == 0 and re_.size == 0:
            return False, False
        return (not _binary_equal(lfs, lp, rfs, rp, ctx)), False
    return (not _rules_equal(lfs, lp, rfs, rp, opts, ctx)), False       # rules


def _time_status(le, re_, tol):
    d = le.mtime - re_.mtime
    if d > tol:
        return "newer"
    if d < -tol:
        return "older"
    return "changed"


def _apply_aligns(L, R, opts, rel):
    """Re-key right entries that the user paired with a differently named left entry."""
    renames = {}
    for lr, rr in opts.aligns:
        lparent, _, lname = lr.rpartition("/")
        rparent, _, rname = rr.rpartition("/")
        if lparent != rel or rparent != rel:
            continue
        lk = lname.lower() if opts.ignore_case_names else lname
        rk = rname.lower() if opts.ignore_case_names else rname
        if lk == rk or lk not in L or rk not in R:
            continue
        ent = R.pop(rk)
        if lk in R:               # a right file already carries the left name: it stays an orphan
            R[lk + "\0"] = R[lk]
        R[lk] = ent
        renames[lk] = rname
    return renames


def compare_tree(lfs, lroot, rfs, rroot, opts=None, rel="", depth=0, counter=None, ctx=None, defer=None):
    """Compare two directory trees (any VFS). Returns a list of Node (dirs first).

    With `defer` (a list) content comparisons are not run here: (node, args) tuples are appended instead and
    resolved later by resolve_deferred() - see compare_tree_parallel()."""
    opts = opts or CompareOptions()
    L = _list(lfs, lroot, opts)
    R = _list(rfs, rroot, opts)
    renames = _apply_aligns(L, R, opts, rel) if opts.aligns else {}
    nodes = []
    for key in sorted(set(L) | set(R), key=lambda k: (not ((L.get(k) or R.get(k))[0].is_dir), k.lower())):
        if ctx is not None:
            if ctx.cancel:
                raise Cancelled()
            ctx.count += 1
        l, r = L.get(key), R.get(key)
        le, lp = l if l else (None, None)
        re_, rp = r if r else (None, None)
        ref = le or re_
        name = (le or re_).name
        is_dir = ref.is_dir
        mismatch = bool(le and re_ and le.is_dir != re_.is_dir)
        sub_l = sub_r = None
        is_arch = False
        if (opts.archives and le and re_ and not le.is_dir and not re_.is_dir
                and vfs.is_archive_name(le.name) and vfs.is_archive_name(re_.name)):
            try:
                sub_l = vfs.ArchiveFS(lfs.local_path(lp))
                sub_r = vfs.ArchiveFS(rfs.local_path(rp))
                is_arch = is_dir = True
            except (OSError, vfs.VFSError):
                sub_l = sub_r = None
        n = Node(name=name, rel=(rel + "/" + name) if rel else name, is_dir=is_dir and not mismatch,
                 depth=depth, rname=renames.get(key, ""), archive=is_arch, mismatch=mismatch,
                 lfs=lfs if le else None, rfs=rfs if re_ else None)
        n.left = _info(lfs, lp, le, le.is_dir) if le else None
        n.right = _info(rfs, rp, re_, re_.is_dir) if re_ else None
        if is_arch:
            n.children = compare_tree(sub_l, "", sub_r, "", opts, n.rel, depth + 1, counter, ctx, defer)
            n.status = "same" if all(c.status == "same" for c in n.children) else "changed"
        elif n.is_dir:
            n.children = compare_tree(lfs if le else None, lp if le and le.is_dir else None,
                                      rfs if re_ else None, rp if re_ and re_.is_dir else None,
                                      opts, n.rel, depth + 1, counter, ctx, defer)
            if opts.include and not n.children:
                continue                # nothing under this folder passes the include filter
            if le and not re_:
                n.status = "left"
            elif re_ and not le:
                n.status = "right"
            else:
                n.status = "same" if all(c.status == "same" for c in n.children) else "changed"
        else:
            if le and not re_:
                n.status = "left"
            elif re_ and not le:
                n.status = "right"
            elif mismatch:
                n.status = "changed"
            elif defer is not None and opts.method != "size_time" and lfs.cheap_algo != "git" and rfs.cheap_algo != "git" \
                    and not (lfs.slow or rfs.slow):
                defer.append((n, lfs, lp, le, rfs, rp, re_))
                n.status = "same"                       # placeholder until resolve_deferred()
            else:
                try:
                    differs, approx = files_differ(lfs, lp, le, rfs, rp, re_, opts, ctx)
                except Cancelled:
                    raise
                except (OSError, vfs.VFSError):
                    differs, approx = True, False
                n.approx = approx
                content_addressed = lfs.cheap_algo == "git" or rfs.cheap_algo == "git"     # commit times say nothing here
                n.status = ("changed" if content_addressed else _time_status(le, re_, opts.time_tol)) if differs else "same"
            if counter is not None:
                counter[n.status] = counter.get(n.status, 0) + 1
        nodes.append(n)
    return nodes


def settle_dirs(nodes):
    """Recompute folder statuses from their children (orphan folders keep left/right)."""
    for n in nodes:
        if n.is_dir:
            settle_dirs(n.children)
            if n.status not in ("left", "right"):
                n.status = "same" if all(c.status == "same" for c in n.children) else "changed"


def resolve_deferred(pending, opts, ctx=None, workers=8):
    """Run the deferred content comparisons on a thread pool (file reads release the GIL)."""
    from concurrent.futures import ThreadPoolExecutor

    def one(item):
        n, lfs, lp, le, rfs, rp, re_ = item
        if ctx is not None and ctx.cancel:
            raise Cancelled()
        try:
            differs, approx = files_differ(lfs, lp, le, rfs, rp, re_, opts, ctx)
        except Cancelled:
            raise
        except (OSError, vfs.VFSError):
            differs, approx = True, False
        n.approx = approx
        n.status = _time_status(le, re_, opts.time_tol) if differs else "same"
        if ctx is not None:
            ctx.count += 1
    if len(pending) < 16:
        for item in pending:
            one(item)
        return
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for _ in ex.map(one, pending):
            pass


def compare_tree_parallel(lfs, lroot, rfs, rroot, opts=None, counter=None, ctx=None, workers=8):
    """compare_tree() with content comparisons spread over several threads. Same result, much faster on big trees."""
    opts = opts or CompareOptions()
    pending = []
    nodes = compare_tree(lfs, lroot, rfs, rroot, opts, counter=None, ctx=ctx, defer=pending)
    if pending:
        resolve_deferred(pending, opts, ctx, workers)
        settle_dirs(nodes)
    if counter is not None:
        counter.clear()
        counter.update(count_statuses(nodes))
    return nodes


def compare_folders(left, right, patterns, rel="", depth=0, counter=None, ctx=None, opts=None):
    """Convenience wrapper for two local folders (kept for scripts and tests)."""
    opts = opts or CompareOptions(patterns=list(patterns))
    lfs = vfs.LocalFS(left) if left else None
    rfs = vfs.LocalFS(right) if right else None
    return compare_tree(lfs, left, rfs, right, opts, rel, depth, counter, ctx)


def compress(nodes, depth=0):
    """Collapse chains of single-child folders (a/b/c) into one row and fix depths."""
    for n in nodes:
        n.depth = depth
        while n.is_dir and not n.archive and len(n.children) == 1 and n.children[0].is_dir and not n.children[0].archive:
            c = n.children[0]
            n.name += "/" + c.name
            if n.rname or c.rname:
                n.rname = (n.rname or n.name.split("/")[0]) + "/" + (c.rname or c.name)
            n.rel, n.children, n.status = c.rel, c.children, c.status
            n.left, n.right = c.left, c.right
            n.lfs, n.rfs = c.lfs, c.rfs
        compress(n.children, depth + 1)
    return nodes


def walk(nodes):
    for n in nodes:
        yield n
        yield from walk(n.children)


def count_statuses(nodes):
    c = {}
    for n in walk(nodes):
        if not n.is_dir:
            c[n.status] = c.get(n.status, 0) + 1
    return c


def flatten(nodes):
    """All file nodes in one list (BC's 'flatten folders'), ordered by relative path."""
    return sorted((n for n in walk(nodes) if not n.is_dir), key=lambda n: n.rel.lower())


def filter_nodes(nodes, keep):
    """Copy of the tree keeping leaf nodes where keep(node) is true (and folders that still have content)."""
    out = []
    for n in nodes:
        if n.is_dir:
            kids = filter_nodes(n.children, keep)
            if kids:
                c = Node(**{**n.__dict__, "children": kids})
                out.append(c)
        elif keep(n):
            out.append(n)
    return out


# =============================================================== text IO
ENCODINGS = {
    # label -> (python codec for decoding, BOM bytes written, codec for encoding)
    "UTF-8": ("utf-8", b"", "utf-8"),
    "UTF-8 BOM": ("utf-8-sig", b"\xef\xbb\xbf", "utf-8"),
    "UTF-16 LE": ("utf-16", b"\xff\xfe", "utf-16-le"),
    "UTF-16 BE": ("utf-16", b"\xfe\xff", "utf-16-be"),
    "UTF-16 LE (no BOM)": ("utf-16-le", b"", "utf-16-le"),
    "UTF-16 BE (no BOM)": ("utf-16-be", b"", "utf-16-be"),
    "UTF-32 LE": ("utf-32", b"\xff\xfe\x00\x00", "utf-32-le"),
    "UTF-32 BE": ("utf-32", b"\x00\x00\xfe\xff", "utf-32-be"),
    "ANSI": ("cp1252", b"", "cp1252"),
    "Latin-1": ("latin-1", b"", "latin-1"),
    "Windows-1250": ("cp1250", b"", "cp1250"),
    "Windows-1251": ("cp1251", b"", "cp1251"),
    "KOI8-R": ("koi8-r", b"", "koi8-r"),
    "Shift-JIS": ("shift_jis", b"", "shift_jis"),
    "GBK": ("gbk", b"", "gbk"),
    "Big5": ("big5", b"", "big5"),
    "EUC-KR": ("euc-kr", b"", "euc-kr"),
}
ENCODING_ALIASES = {"UTF-16": "UTF-16 LE", "UTF-32": "UTF-32 LE"}


def detect_encoding(raw):
    """Pick an encoding label from BOMs and byte patterns."""
    if raw.startswith(b"\xff\xfe\x00\x00"):
        return "UTF-32 LE"
    if raw.startswith(b"\x00\x00\xfe\xff"):
        return "UTF-32 BE"
    if raw.startswith(b"\xef\xbb\xbf"):
        return "UTF-8 BOM"
    if raw.startswith(b"\xff\xfe"):
        return "UTF-16 LE"
    if raw.startswith(b"\xfe\xff"):
        return "UTF-16 BE"
    sample = raw[:8192]
    if len(sample) >= 4 and b"\x00" in sample:
        even = sample[0::2].count(0)
        odd = sample[1::2].count(0)
        half = max(1, len(sample) // 2)
        if odd / half > 0.3 and even / half < 0.05:
            return "UTF-16 LE (no BOM)"
        if even / half > 0.3 and odd / half < 0.05:
            return "UTF-16 BE (no BOM)"
    try:
        raw.decode("utf-8")
        return "UTF-8"
    except UnicodeDecodeError as e:
        # a multi-byte sequence cut off at the end of a very large buffer is still UTF-8
        if e.start >= len(raw) - 3 and len(raw) > 65536:
            return "UTF-8"
    return "ANSI"


def detect_eol(text):
    """Return (dominant label, is_mixed). Labels: CRLF, LF, CR."""
    crlf = text.count("\r\n")
    cr = text.count("\r") - crlf
    lf = text.count("\n") - crlf
    kinds = [(crlf, "CRLF"), (lf, "LF"), (cr, "CR")]
    present = [k for k in kinds if k[0] > 0]
    if not present:
        return "LF", False
    present.sort(reverse=True)
    return present[0][1], len(present) > 1


def decode_bytes(raw, encoding=None):
    """Return (text, encoding label, eol label). `encoding` forces a label."""
    label = ENCODING_ALIASES.get(encoding, encoding) if encoding else detect_encoding(raw)
    if label not in ENCODINGS:
        label = "UTF-8"
    codec = ENCODINGS[label][0]
    try:
        text = raw.decode(codec)
    except UnicodeDecodeError:
        text = raw.decode(codec, errors="replace")
    eol, _ = detect_eol(text)
    return text, label, eol


def read_text(path, encoding=None):
    """Return (text, encoding label, eol label)."""
    with open(path, "rb") as f:
        raw = f.read()
    return decode_bytes(raw, encoding)


@dataclass
class TextFile:
    lines: list
    encoding: str
    eol: str
    mixed_eol: bool
    mtime: float
    final_newline: bool = True


def load_text_file(path, encoding=None, raw=None):
    if raw is None:
        with open(path, "rb") as f:
            raw = f.read()
    text, enc, eol = decode_bytes(raw, encoding)
    mixed = detect_eol(text)[1]
    try:
        mt = os.path.getmtime(path)
    except OSError:
        mt = None
    return TextFile(split_lines(text), enc, eol, mixed, mt, final_newline=(not text) or text.endswith(("\n", "\r")))


def encode_text(text, label):
    label = ENCODING_ALIASES.get(label, label)
    codec_dec, bom, codec_enc = ENCODINGS.get(label, ENCODINGS["UTF-8"])
    return bom + text.encode(codec_enc, errors="replace")


def write_text(path, lines, encoding_label, eol, final_newline=True):
    nl = {"CRLF": "\r\n", "CR": "\r"}.get(eol, "\n")
    text = nl.join(lines)
    if lines and final_newline:
        text += nl
    data = encode_text(text, encoding_label)
    tmp = path + ".lavati-tmp"
    with open(tmp, "wb") as f:
        f.write(data)
    try:
        st = os.stat(path)
        os.chmod(tmp, st.st_mode)
    except OSError:
        pass
    os.replace(tmp, path)


_SPLIT = re.compile(r"\r\n|\r|\n")


def split_lines(text):
    if not text:
        return []
    parts = _SPLIT.split(text)
    if parts and parts[-1] == "":
        parts.pop()
    return parts


def line_key(s, ignore_ws, ignore_case=False):
    if ignore_ws:
        s = " ".join(s.split())
    if ignore_case:
        s = s.lower()
    return s


# =============================================================== text diff
def _sm_opcodes(ka, kb):
    """difflib opcodes with common prefix/suffix trimmed first (big speed-up for small edits)."""
    n, m = len(ka), len(kb)
    pre = 0
    while pre < n and pre < m and ka[pre] == kb[pre]:
        pre += 1
    suf = 0
    while suf < n - pre and suf < m - pre and ka[n - 1 - suf] == kb[m - 1 - suf]:
        suf += 1
    ops = []
    if pre:
        ops.append(("equal", 0, pre, 0, pre))
    ma, mb = ka[pre:n - suf], kb[pre:m - suf]
    if ma or mb:
        if not ma:
            ops.append(("insert", pre, pre, pre, pre + len(mb)))
        elif not mb:
            ops.append(("delete", pre, pre + len(ma), pre, pre))
        else:
            sm = difflib.SequenceMatcher(None, ma, mb, autojunk=False)
            for tag, i1, i2, j1, j2 in sm.get_opcodes():
                ops.append((tag, i1 + pre, i2 + pre, j1 + pre, j2 + pre))
    if suf:
        ops.append(("equal", n - suf, n, m - suf, m))
    return ops


def diff_opcodes(a, b, ignore_ws=False, ignore_case=False, rules=None, anchors=None, eol_tags=("", "")):
    """Opcodes (difflib style, incl. 'equal') comparing line lists a and b.

    rules   - TextRules; overrides ignore_ws / ignore_case when given
    anchors - [(i, j), ...] line pairs forced to align (strictly increasing)
    """
    if rules is None:
        rules = TextRules(ws="changes" if ignore_ws else "none", ignore_case=ignore_case)
    ka, kb = rules.keys(a, eol_tags[0]), rules.keys(b, eol_tags[1])
    if not anchors:
        return _sm_opcodes(ka, kb)
    ops = []
    pi = pj = 0
    clean = []
    last = (-1, -1)
    for i, j in sorted(anchors):
        if 0 <= i < len(a) and 0 <= j < len(b) and i > last[0] and j > last[1]:
            clean.append((i, j))
            last = (i, j)
    for i, j in clean + [(len(a), len(b))]:
        for tag, i1, i2, j1, j2 in _sm_opcodes(ka[pi:i], kb[pj:j]):
            ops.append((tag, i1 + pi, i2 + pi, j1 + pj, j2 + pj))
        if i < len(a) and j < len(b):
            tag = "equal" if ka[i] == kb[j] else "replace"
            ops.append((tag, i, i + 1, j, j + 1))
        pi, pj = i + 1, j + 1
    # merge neighbouring equal blocks
    merged = []
    for op in ops:
        if merged and op[0] == "equal" and merged[-1][0] == "equal" and merged[-1][2] == op[1] and merged[-1][4] == op[3]:
            t, i1, _, j1, _ = merged[-1]
            merged[-1] = (t, i1, op[2], j1, op[4])
        else:
            merged.append(op)
    return merged


_TOKEN = re.compile(r"\w+|\s+|[^\w\s]")


def _tokens(s, rules):
    """[(key, start, end)] for intra-line diffing, honouring whitespace / case rules."""
    out = []
    if rules is not None and rules.ws == "all":      # whitespace is invisible: diff character by character
        for i, ch in enumerate(s):
            if not ch.isspace():
                out.append((ch.lower() if rules.ignore_case else ch, i, i + 1))
        return out
    for m in _TOKEN.finditer(s):
        t = m.group()
        if t.isspace():
            if rules is not None and rules.ws in ("edges", "changes"):
                if rules.ws == "edges" and (m.start() == 0 or m.end() == len(s)):
                    continue
                t = " "
        if rules is not None and rules.ignore_case:
            t = t.lower()
        out.append((t, m.start(), m.end()))
    return out


def _char_refine(a, b, a0, b0, sa, sb):
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    ops = sm.get_opcodes()
    # tiny equal islands inside a changed word make the highlight look noisy: merge them
    spans = []
    for tag, i1, i2, j1, j2 in ops:
        if tag != "equal":
            spans.append([i1, i2, j1, j2])
        elif spans and (i2 - i1) < 2 and tag == "equal":
            spans.append([i1, i2, j1, j2, "gap"])
    merged = []
    for s in spans:
        if len(s) == 5:
            merged.append(s)
            continue
        if len(merged) >= 2 and len(merged[-1]) == 5 and len(merged[-2]) == 4:
            gap = merged.pop()
            prev = merged.pop()
            s = [prev[0], s[1], prev[2], s[3]]
        merged.append(s)
    for s in merged:
        if len(s) == 5:
            continue
        if s[1] > s[0]:
            sa.append((a0 + s[0], a0 + s[1]))
        if s[3] > s[2]:
            sb.append((b0 + s[2], b0 + s[3]))


def word_spans(a, b, rules=None, chars=False):
    """Return (spans_a, spans_b): lists of (start, end) char ranges that differ.

    chars=True refines single changed words down to the differing characters.
    """
    ta, tb = _tokens(a, rules), _tokens(b, rules)
    sm = difflib.SequenceMatcher(None, [t[0] for t in ta], [t[0] for t in tb], autojunk=False)
    sa, sb = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if chars and tag == "replace" and i2 - i1 == 1 and j2 - j1 == 1:
            x, y = ta[i1], tb[j1]
            r = difflib.SequenceMatcher(None, x[0], y[0], autojunk=False).ratio()
            if r >= 0.4 and len(x[0]) > 1 and len(y[0]) > 1:
                _char_refine(a[x[1]:x[2]], b[y[1]:y[2]], x[1], y[1], sa, sb)
                continue
        if i2 > i1:
            sa.append((ta[i1][1], ta[i2 - 1][2]))
        if j2 > j1:
            sb.append((tb[j1][1], tb[j2 - 1][2]))
    return sa, sb


def group_opcodes(ops, n=3):
    """Like difflib.SequenceMatcher.get_grouped_opcodes, for an existing opcode list."""
    codes = list(ops) or [("equal", 0, 1, 0, 1)]
    if codes[0][0] == "equal":
        t, i1, i2, j1, j2 = codes[0]
        codes[0] = t, max(i1, i2 - n), i2, max(j1, j2 - n), j2
    if codes[-1][0] == "equal":
        t, i1, i2, j1, j2 = codes[-1]
        codes[-1] = t, i1, min(i2, i1 + n), j1, min(j2, j1 + n)
    nn = n + n
    group = []
    for tag, i1, i2, j1, j2 in codes:
        if tag == "equal" and i2 - i1 > nn:
            group.append((tag, i1, min(i2, i1 + n), j1, min(j2, j1 + n)))
            yield group
            group = []
            i1, j1 = max(i1, i2 - n), max(j1, j2 - n)
        group.append((tag, i1, i2, j1, j2))
    if group and not (len(group) == 1 and group[0][0] == "equal"):
        yield group


def unified_diff(a, b, left_name="left", right_name="right", context=3, rules=None, anchors=None,
                 left_time="", right_time=""):
    """Unified-diff text (list of lines, no trailing newlines) for two line lists."""
    ops = diff_opcodes(a, b, rules=rules or TextRules(), anchors=anchors)
    out = []
    started = False
    for group in group_opcodes(ops, context):
        if not started:
            out.append(f"--- {left_name}" + (f"\t{left_time}" if left_time else ""))
            out.append(f"+++ {right_name}" + (f"\t{right_time}" if right_time else ""))
            started = True
        first, last = group[0], group[-1]

        def rng(s, e):
            ln = e - s
            if ln == 1:
                return f"{s + 1}"
            return f"{(s + 1) if ln else s},{ln}"
        out.append(f"@@ -{rng(first[1], last[2])} +{rng(first[3], last[4])} @@")
        for tag, i1, i2, j1, j2 in group:
            if tag == "equal":
                out += [" " + x for x in a[i1:i2]]
                continue
            if tag in ("replace", "delete"):
                out += ["-" + x for x in a[i1:i2]]
            if tag in ("replace", "insert"):
                out += ["+" + x for x in b[j1:j2]]
    return out


def text_pairs(L, R, ops, rules):
    """Aligned row model shared by the views and reports.

    Returns (pairs, hunks): pairs = [(li, ri, kind, hunk)] with kind eq|unimp|mod|del|ins.
    """
    pairs, hunks = [], []
    for tag, i1, i2, j1, j2 in ops:
        if tag == "equal":
            for k in range(i2 - i1):
                li, ri = i1 + k, j1 + k
                pairs.append((li, ri, "eq" if L[li] == R[ri] else "unimp", -1))
            continue
        h = len(hunks)
        hunks.append((i1, i2, j1, j2))
        nl, nr = i2 - i1, j2 - j1
        m = min(nl, nr)
        for k in range(m):
            pairs.append((i1 + k, j1 + k, "mod", h))
        for k in range(m, nl):
            pairs.append((i1 + k, None, "del", h))
        for k in range(m, nr):
            pairs.append((None, j1 + k, "ins", h))
    return pairs, hunks


# =============================================================== merges
@dataclass
class Segment:
    kind: str                 # "text" | "auto" | "conflict"
    lines: list = field(default_factory=list)   # resolved lines (text/auto)
    ours: list = field(default_factory=list)
    base: list = field(default_factory=list)
    theirs: list = field(default_factory=list)
    side: str = ""            # for auto: "ours"/"theirs"/"both"
    resolution: str = None    # for conflict: ours/theirs/both/base/manual ; for auto: override
    manual: list = None
    ours_range: tuple = (0, 0)
    theirs_range: tuple = (0, 0)
    base_range: tuple = (0, 0)


def merge2(left, right, rules=None):
    """Two-way merge: every difference is a choice between left and right."""
    ops = diff_opcodes(left, right, rules=rules or TextRules())
    segs = []
    for tag, i1, i2, j1, j2 in ops:
        if tag == "equal":
            segs.append(Segment("text", left[i1:i2]))
        else:
            s = Segment("conflict", ours=left[i1:i2], theirs=right[j1:j2], base=[])
            s.ours_range, s.theirs_range, s.base_range = (i1, i2), (j1, j2), (0, 0)
            segs.append(s)
    return segs


def merge3(ours, base, theirs, rules=None):
    rules = rules or TextRules()

    def changes(other):
        ka, kb = rules.keys(base), rules.keys(other)
        sm = difflib.SequenceMatcher(None, ka, kb, autojunk=False)
        return [(i1, i2, other[j1:j2], (j1, j2)) for tag, i1, i2, j1, j2 in sm.get_opcodes() if tag != "equal"]

    co = [("ours",) + c for c in changes(ours)]
    ct = [("theirs",) + c for c in changes(theirs)]
    allc = sorted(co + ct, key=lambda c: (c[1], c[2]))
    clusters = []
    for c in allc:
        if clusters and (c[1] < clusters[-1]["end"] or
                         (c[1] == clusters[-1]["end"] == clusters[-1]["start"])):
            cl = clusters[-1]
            cl["end"] = max(cl["end"], c[2])
            cl["items"].append(c)
        else:
            clusters.append({"start": c[1], "end": c[2], "items": [c]})

    def apply(side_items, s, e):
        out, pos = [], s
        for _, i1, i2, rep, _rng in side_items:
            out += base[pos:i1]
            out += rep
            pos = i2
        out += base[pos:e]
        return out

    segs = []
    pos = 0
    for cl in clusters:
        s, e = cl["start"], cl["end"]
        if s > pos:
            segs.append(Segment("text", base[pos:s]))
        io = [c for c in cl["items"] if c[0] == "ours"]
        it = [c for c in cl["items"] if c[0] == "theirs"]
        o_lines = apply(io, s, e)
        t_lines = apply(it, s, e)
        b_lines = base[s:e]

        def rng(items):
            if not items:
                return None
            return (min(c[4][0] for c in items), max(c[4][1] for c in items))
        if io and not it:
            seg = Segment("auto", o_lines, side="ours")
        elif it and not io:
            seg = Segment("auto", t_lines, side="theirs")
        elif rules.keys(o_lines) == rules.keys(t_lines):
            seg = Segment("auto", o_lines, side="both")
        else:
            seg = Segment("conflict", ours=o_lines, base=b_lines, theirs=t_lines)
        seg.ours, seg.theirs, seg.base = o_lines, t_lines, b_lines
        seg.ours_range = rng(io) or (0, 0)
        seg.theirs_range = rng(it) or (0, 0)
        seg.base_range = (s, e)
        segs.append(seg)
        pos = e
    if pos < len(base):
        segs.append(Segment("text", base[pos:]))
    return segs


def segment_output(seg):
    if seg.kind == "text":
        return seg.lines
    r = seg.resolution
    if seg.kind == "auto":
        if r in ("base", "ours", "theirs") and r != seg.side:
            return {"base": seg.base, "ours": seg.ours, "theirs": seg.theirs}[r]
        return seg.lines
    if r == "ours":
        return seg.ours
    if r == "theirs":
        return seg.theirs
    if r == "both":
        return seg.ours + seg.theirs
    if r == "both_rev":
        return seg.theirs + seg.ours
    if r == "base":
        return seg.base
    if r == "manual":
        return seg.manual or []
    return None


# =============================================================== folder merge
def classify3(ours, base, theirs):
    """File-level three-way classification from (exists, crc/identity) tuples.

    Each argument is None (absent) or a comparable token (e.g. content hash).
    Returns (action, why): action in keep|ours|theirs|base|delete|conflict.
    """
    if ours == theirs:
        return ("keep" if ours is not None else "delete"), "same on both sides"
    if ours == base:
        return ("theirs" if theirs is not None else "delete"), "changed only in theirs"
    if theirs == base:
        return ("ours" if ours is not None else "delete"), "changed only in ours"
    return "conflict", "changed on both sides"


# =============================================================== tables
def list_table_sources(path):
    """Sheet names (Excel) or table labels (HTML). Empty list for CSV/TSV."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    if ext in HTML_EXTS:
        tables = _html_tables(read_text(path)[0])
        return [f"Table {i + 1} ({len(t)} rows)" for i, t in enumerate(tables)]
    return []


def _html_tables(text):
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.tables, self.stack = [], []
            self.row = self.cell = None

        def handle_starttag(self, tag, attrs):
            if tag == "table":
                self.stack.append([])
            elif tag == "tr" and self.stack:
                self.row = []
            elif tag in ("td", "th") and self.stack is not None and self.row is not None:
                self.cell = []

        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.cell is not None and self.row is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.cell = None
            elif tag == "tr" and self.row is not None and self.stack:
                if self.row:
                    self.stack[-1].append(self.row)
                self.row = None
            elif tag == "table" and self.stack:
                self.tables.append(self.stack.pop())

        def handle_data(self, data):
            if self.cell is not None:
                self.cell.append(data)
    p = P()
    p.feed(text)
    p.close()
    return [t for t in p.tables if t]


def load_table(path, sheet=None):
    """Return (headers, body). `sheet` = sheet name (Excel) or table index/label (HTML)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet in wb.sheetnames else wb.active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext == ".xls":
        raise ValueError("Legacy .xls files are not supported - save as .xlsx or .csv first")
    elif ext in HTML_EXTS:
        tables = _html_tables(read_text(path)[0])
        if not tables:
            raise ValueError("No <table> found in this HTML file")
        idx = 0
        if isinstance(sheet, int):
            idx = sheet
        elif isinstance(sheet, str) and sheet.startswith("Table "):
            try:
                idx = int(sheet.split()[1]) - 1
            except (ValueError, IndexError):
                idx = 0
        rows = tables[min(idx, len(tables) - 1)]
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
    else:
        text, _, _ = read_text(path)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        rows = list(csv.reader(text.splitlines(), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)] or [[]]
    headers = [h.strip() or f"col{i + 1}" for i, h in enumerate(rows[0])]
    seen = {}
    for i, h in enumerate(headers):       # duplicate header names would collide as dict keys
        if h in seen:
            seen[h] += 1
            headers[i] = f"{h}.{seen[h]}"
        else:
            seen[h] = 0
    body = [r + [""] * (len(headers) - len(r)) for r in rows[1:]]
    body = [r[:len(headers)] if len(r) > len(headers) else r for r in body]
    return headers, body


@dataclass
class TableResult:
    headers: list
    rows: list          # dict(kind, key, cells=[(left,right,changed)])
    counts: dict


def _norm(v, ignore_case, trim):
    if trim:
        v = v.strip()
    if ignore_case:
        v = v.lower()
    return v


def _cell_equal(a, b, ignore_case, trim, num_as_text, tol, relative=False):
    if _norm(a, ignore_case, trim) == _norm(b, ignore_case, trim):
        return True
    if not num_as_text:
        try:
            fa = float(a.replace(",", "").strip())
            fb = float(b.replace(",", "").strip())
            if relative:                       # tol is a percentage of the larger magnitude
                scale = max(abs(fa), abs(fb))
                return abs(fa - fb) <= scale * tol / 100.0 + 1e-12
            return abs(fa - fb) <= tol + 1e-12
        except ValueError:
            pass
    return False


def compare_tables(L, R, key, compare_cols, ignore_case=True, trim=True, num_as_text=False, tol=0.0, relative=False):
    lh, lrows = L
    rh, rrows = R
    headers = list(lh) + [h for h in rh if h not in lh]

    def index(h, rows):
        pos = {name: i for i, name in enumerate(h)}
        out = {}
        order = []
        seen = {}
        for r in rows:
            k = _norm(r[pos[key]], ignore_case, trim) if key in pos else ""
            n = seen.get(k, 0)
            seen[k] = n + 1
            kk = (k, n)
            out[kk] = {name: r[i] for name, i in pos.items()}
            order.append(kk)
        return out, order

    li, lorder = index(lh, lrows)
    ri, rorder = index(rh, rrows)
    result, counts = [], {"changed": 0, "left": 0, "right": 0, "same": 0}
    keys = list(lorder) + [k for k in rorder if k not in li]
    for k in keys:
        a, b = li.get(k), ri.get(k)
        cells = []
        if a and b:
            changed_any = False
            for h in headers:
                va, vb = a.get(h, ""), b.get(h, "")
                ch = h in compare_cols and not _cell_equal(va, vb, ignore_case, trim, num_as_text, tol, relative)
                changed_any |= ch
                cells.append((va, vb, ch))
            kind = "changed" if changed_any else "same"
        elif a:
            kind = "left"
            cells = [(a.get(h, ""), "", False) for h in headers]
        else:
            kind = "right"
            cells = [("", b.get(h, ""), False) for h in headers]
        counts[kind] += 1
        keyval = (a or b).get(key, "")
        result.append({"kind": kind, "key": keyval, "cells": cells})
    return TableResult(headers, result, counts)
